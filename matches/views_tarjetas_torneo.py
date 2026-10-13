from decimal import Decimal
import io
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.utils import timezone
from django.db import models, transaction
from django.http import HttpResponse

from matches.models import Torneo, GrupoTorneo, Partido, EventoPartido
from teams.models import Equipo
from finances.models import MultaTarjeta, MovimientoCaja


def _get_torneo_safe(request, torneo_id):
    if hasattr(request, 'organizacion') and request.organizacion:
        return get_object_or_404(Torneo, id=torneo_id, organizacion=request.organizacion)
    return get_object_or_404(Torneo, id=torneo_id)


def _obtener_mapa_grupos_torneo(torneo):
    """
    Retorna:
    - grupos_equipos: dict {equipo_id: nombre_grupo}
    - lista_grupos: list de dicts con {'id': id_or_str, 'nombre': str, 'equipos_ids': [ids]}
    """
    grupos_equipos = {}
    lista_grupos = []

    # 1. Si tiene grupos personalizados (Fase 2 Personalizado)
    try:
        grupos_p = list(torneo.grupos_personalizados.prefetch_related('equipos_asignados__equipo').all())
    except Exception:
        grupos_p = []

    if grupos_p:
        for g in grupos_p:
            eq_ids = []
            for ea in g.equipos_asignados.all():
                grupos_equipos[ea.equipo_id] = g.nombre
                eq_ids.append(ea.equipo_id)
            lista_grupos.append({
                'id': g.id,
                'nombre': g.nombre,
                'equipos_ids': eq_ids
            })
        return grupos_equipos, lista_grupos

    # 2. Si usa distribucion_grupos (Torneo tradicional grupos)
    dist = getattr(torneo, 'distribucion_grupos', None)
    if isinstance(dist, dict) and dist:
        for g_nombre, eq_list in dist.items():
            eq_ids = []
            for item in eq_list:
                try:
                    eid = int(item)
                    grupos_equipos[eid] = g_nombre
                    eq_ids.append(eid)
                except (ValueError, TypeError):
                    pass
            lista_grupos.append({
                'id': g_nombre,
                'nombre': g_nombre,
                'equipos_ids': eq_ids
            })
        return grupos_equipos, lista_grupos

    # 3. Torneo único / liga sin grupos
    all_eq_ids = [eq.id for eq in torneo.equipos.all()]
    for eid in all_eq_ids:
        grupos_equipos[eid] = 'General'
    lista_grupos.append({
        'id': 'general',
        'nombre': 'General',
        'equipos_ids': all_eq_ids
    })
    return grupos_equipos, lista_grupos


def _calcular_datos_tarjetas(torneo, org):
    """
    Calcula de manera centralizada y optimizada todos los datos disciplinarios y financieros
    del torneo: KPIs, agrupación por equipo con detalle de multas, agrupación por fecha/jornada
    con partidos y eventos, agrupación por grupos e historial completo.
    """
    # 1. Sincronización proactiva: asegurar que todo evento de tarjeta tenga su MultaTarjeta
    eventos_tarjetas = EventoPartido.objects.filter(
        partido__torneo=torneo,
        tipo__in=['amarilla', 'roja']
    ).select_related('partido', 'jugador', 'equipo')

    costo_amarilla_torneo = Decimal(str(torneo.costo_amarilla if torneo.costo_amarilla is not None else 50.00))
    costo_roja_torneo = Decimal(str(torneo.costo_roja if torneo.costo_roja is not None else 150.00))

    for ev in eventos_tarjetas:
        tipo_str = str(ev.tipo).lower().strip()
        costo = costo_amarilla_torneo if tipo_str == 'amarilla' else costo_roja_torneo

        equipo_dest = ev.equipo
        if not equipo_dest and ev.jugador:
            fichas = getattr(ev.jugador, 'fichas_jugador', None)
            if fichas:
                f = fichas.filter(torneo=torneo).first() or fichas.first()
                if f:
                    equipo_dest = f.equipo
        if not equipo_dest and ev.partido:
            equipo_dest = ev.partido.equipo_local

        if equipo_dest:
            MultaTarjeta.objects.get_or_create(
                evento=ev,
                defaults={
                    'partido': ev.partido,
                    'equipo': equipo_dest,
                    'jugador': ev.jugador,
                    'organizacion': org,
                    'monto': costo,
                    'motivo': tipo_str,
                    'estado': 'pendiente'
                }
            )

    # 2. Obtener todas las multas del torneo
    multas_filter = {'partido__torneo': torneo}
    if org:
        multas_filter['organizacion'] = org

    multas_qs = MultaTarjeta.objects.filter(**multas_filter).select_related(
        'jugador', 'equipo', 'partido', 'evento'
    )

    # 3. KPIs Globales
    total_amarillas = multas_qs.filter(motivo='amarilla').count()
    monto_amarillas = multas_qs.filter(motivo='amarilla').aggregate(models.Sum('monto'))['monto__sum'] or Decimal('0.00')

    total_rojas = multas_qs.filter(motivo='roja').count()
    monto_rojas = multas_qs.filter(motivo='roja').aggregate(models.Sum('monto'))['monto__sum'] or Decimal('0.00')

    total_facturado = monto_amarillas + monto_rojas
    total_pagado = multas_qs.filter(estado='pagado').aggregate(models.Sum('monto'))['monto__sum'] or Decimal('0.00')
    total_pendiente = total_facturado - total_pagado

    totales = {
        'amarillas_count': total_amarillas,
        'rojas_count': total_rojas,
        'total_sanciones': total_amarillas + total_rojas,
        'monto_amarillas': monto_amarillas,
        'monto_rojas': monto_rojas,
        'monto_total': total_facturado,
        'monto_pagado': total_pagado,
        'monto_pendiente': total_pendiente,
        'total_multas_count': multas_qs.count(),
        'multas_pagadas_count': multas_qs.filter(estado='pagado').count(),
        'multas_pendientes_count': multas_qs.filter(estado='pendiente').count(),
    }

    # 4. Agrupación por Equipo
    grupos_equipos, lista_grupos = _obtener_mapa_grupos_torneo(torneo)

    equipos_dict = {}
    for eq in torneo.equipos.all():
        equipos_dict[eq.id] = {
            'equipo': eq,
            'grupo_nombre': grupos_equipos.get(eq.id, "General"),
            'amarillas': 0,
            'rojas': 0,
            'total_tarjetas': 0,
            'monto_total': Decimal('0.00'),
            'monto_pagado': Decimal('0.00'),
            'monto_pendiente': Decimal('0.00'),
            'tiene_deuda': False,
            'multas_pendientes_ids': [],
            'multas_list': [],
        }

    for m in multas_qs:
        if m.equipo_id not in equipos_dict:
            equipos_dict[m.equipo_id] = {
                'equipo': m.equipo,
                'grupo_nombre': grupos_equipos.get(m.equipo_id, "General"),
                'amarillas': 0,
                'rojas': 0,
                'total_tarjetas': 0,
                'monto_total': Decimal('0.00'),
                'monto_pagado': Decimal('0.00'),
                'monto_pendiente': Decimal('0.00'),
                'tiene_deuda': False,
                'multas_pendientes_ids': [],
                'multas_list': [],
            }
        ed = equipos_dict[m.equipo_id]
        ed['total_tarjetas'] += 1
        ed['multas_list'].append(m)
        if m.motivo == 'amarilla':
            ed['amarillas'] += 1
        elif m.motivo == 'roja':
            ed['rojas'] += 1
        ed['monto_total'] += m.monto
        if m.estado == 'pagado':
            ed['monto_pagado'] += m.monto
        else:
            ed['monto_pendiente'] += m.monto
            ed['tiene_deuda'] = True
            ed['multas_pendientes_ids'].append(m.id)

    equipos_data = sorted(
        equipos_dict.values(),
        key=lambda x: (-x['monto_pendiente'], -x['total_tarjetas'], x['equipo'].nombre if x['equipo'] else '')
    )

    # 5. Agrupación por Fecha / Jornada con partidos y eventos
    multas_por_evento = {m.evento_id: m for m in multas_qs if m.evento_id}
    partidos_torneo = Partido.objects.filter(torneo=torneo).select_related(
        'equipo_local', 'equipo_visitante'
    ).prefetch_related('eventos', 'eventos__jugador', 'eventos__equipo').order_by('jornada', 'fecha_hora')

    fechas_dict = {}
    for p in partidos_torneo:
        f_num = p.jornada if p.jornada else 1
        if f_num not in fechas_dict:
            fechas_dict[f_num] = {
                'numero_fecha': f_num,
                'total_amarillas': 0,
                'total_rojas': 0,
                'total_multas_monto': Decimal('0.00'),
                'total_pagado_monto': Decimal('0.00'),
                'total_pendiente_monto': Decimal('0.00'),
                'partidos': []
            }
        
        fd = fechas_dict[f_num]
        
        p_amarillas = 0
        p_rojas = 0
        p_monto = Decimal('0.00')
        p_pagado = Decimal('0.00')
        p_eventos = []

        for ev in p.eventos.all():
            if ev.tipo in ['amarilla', 'roja']:
                m = multas_por_evento.get(ev.id)
                costo_ev = m.monto if m else (costo_amarilla_torneo if ev.tipo == 'amarilla' else costo_roja_torneo)
                pagada_ev = m.estado == 'pagado' if m else False

                if ev.tipo == 'amarilla':
                    p_amarillas += 1
                    fd['total_amarillas'] += 1
                else:
                    p_rojas += 1
                    fd['total_rojas'] += 1

                p_monto += costo_ev
                fd['total_multas_monto'] += costo_ev

                if pagada_ev:
                    p_pagado += costo_ev
                    fd['total_pagado_monto'] += costo_ev
                else:
                    fd['total_pendiente_monto'] += costo_ev

                p_eventos.append({
                    'evento': ev,
                    'multa': m,
                    'monto': costo_ev,
                    'pagada': pagada_ev
                })

        fd['partidos'].append({
            'partido': p,
            'amarillas': p_amarillas,
            'rojas': p_rojas,
            'monto_multas': p_monto,
            'monto_pagado': p_pagado,
            'monto_pendiente': p_monto - p_pagado,
            'eventos': p_eventos
        })

    def _sort_f(item):
        val = str(item['numero_fecha'])
        return (0, int(val)) if val.isdigit() else (1, val)

    fechas_data = sorted(fechas_dict.values(), key=_sort_f)

    # 6. Agrupación por Grupos
    grupos_data = []
    for g_info in lista_grupos:
        g_item = {
            'id': g_info['id'],
            'nombre': g_info['nombre'],
            'total_amarillas': 0,
            'total_rojas': 0,
            'monto_total': Decimal('0.00'),
            'monto_pagado': Decimal('0.00'),
            'monto_pendiente': Decimal('0.00'),
            'equipos': []
        }
        for eq_id in g_info['equipos_ids']:
            if eq_id in equipos_dict:
                ed = equipos_dict[eq_id]
                g_item['equipos'].append(ed)
                g_item['total_amarillas'] += ed['amarillas']
                g_item['total_rojas'] += ed['rojas']
                g_item['monto_total'] += ed['monto_total']
                g_item['monto_pagado'] += ed['monto_pagado']
                g_item['monto_pendiente'] += ed['monto_pendiente']
        grupos_data.append(g_item)

    # 7. Historial cronológico completo de tarjetas
    historial_tarjetas = []
    for m in multas_qs.order_by('-partido__jornada', '-partido__fecha_hora', '-evento__minuto', '-id'):
        historial_tarjetas.append({
            'numero_fecha': m.partido.jornada if (m.partido and m.partido.jornada) else "-",
            'evento': m.evento or m,
            'multa': m,
            'monto': m.monto,
            'pagada': m.estado == 'pagado'
        })

    lista_equipos = sorted([ed['equipo'] for ed in equipos_data if ed['equipo']], key=lambda x: x.nombre)
    lista_jornadas = [fd['numero_fecha'] for fd in fechas_data]

    return {
        'totales': totales,
        'equipos_data': equipos_data,
        'fechas_data': fechas_data,
        'grupos_data': grupos_data,
        'historial_tarjetas': historial_tarjetas,
        'multas_qs': multas_qs,
        'lista_equipos': lista_equipos,
        'lista_jornadas': lista_jornadas,
        'grupos_equipos': grupos_equipos,
    }


@login_required
def control_tarjetas_torneo_view(request, torneo_id):
    """
    Vista integral de Control de Tarjetas (Amarillas y Rojas) y Multas Financieras del Torneo.
    Permite visualizar y calcular por Fecha/Jornada, por Equipo y por Grupos, con enlace directo
    a Tesorería para registrar pagos individuales o colectivos.
    """
    torneo = _get_torneo_safe(request, torneo_id)
    org = getattr(request, 'organizacion', None) or torneo.organizacion

    datos = _calcular_datos_tarjetas(torneo, org)

    es_admin = bool(
        request.user.is_superuser or 
        getattr(request.user, 'role', '') in ['tesorero', 'tesoreria', 'superadmin', 'comision', 'organizador']
    )

    context = {
        'torneo': torneo,
        'es_admin': es_admin,
        'totales': datos['totales'],
        'equipos_data': datos['equipos_data'],
        'fechas_data': datos['fechas_data'],
        'grupos_data': datos['grupos_data'],
        'historial_tarjetas': datos['historial_tarjetas'],
    }
    return render(request, 'matches/control_tarjetas_torneo.html', context)


@login_required
def imprimir_reporte_tarjetas_torneo(request, torneo_id):
    """
    Vista especializada de Impresión / Exportación PDF de Sanciones y Multas Financieras.
    Permite generar reportes:
    - total: Consolidado general de todos los equipos del torneo.
    - equipo: Estado de cuenta y liquidación individual de un equipo o de todos.
    - jornada: Planilla disciplinaria por fecha de programación con partidos y amonestados.
    - historial: Registro cronológico completo de amonestaciones.
    """
    torneo = _get_torneo_safe(request, torneo_id)
    org = getattr(request, 'organizacion', None) or torneo.organizacion

    datos = _calcular_datos_tarjetas(torneo, org)

    tipo = request.GET.get('tipo', 'total')  # 'total', 'equipo', 'jornada', 'historial'
    equipo_id_raw = request.GET.get('equipo_id', '')
    jornada_raw = request.GET.get('jornada', '')
    filtro_estado = request.GET.get('estado', 'todos')  # 'todos', 'pendiente', 'pagado'

    equipo_seleccionado = None
    if equipo_id_raw and equipo_id_raw.isdigit():
        eid = int(equipo_id_raw)
        for ed in datos['equipos_data']:
            if ed['equipo'] and ed['equipo'].id == eid:
                equipo_seleccionado = ed
                break

    jornada_seleccionada = None
    if jornada_raw and jornada_raw.isdigit():
        jnum = int(jornada_raw)
        for fd in datos['fechas_data']:
            if fd['numero_fecha'] == jnum:
                jornada_seleccionada = fd
                break

    context = {
        'torneo': torneo,
        'organizacion': org,
        'totales': datos['totales'],
        'equipos_data': datos['equipos_data'],
        'fechas_data': datos['fechas_data'],
        'grupos_data': datos['grupos_data'],
        'historial_tarjetas': datos['historial_tarjetas'],
        'lista_equipos': datos['lista_equipos'],
        'lista_jornadas': datos['lista_jornadas'],
        'tipo': tipo,
        'equipo_id': equipo_id_raw,
        'equipo_seleccionado': equipo_seleccionado,
        'jornada': jornada_raw,
        'jornada_seleccionada': jornada_seleccionada,
        'filtro_estado': filtro_estado,
        'fecha_impresion': timezone.now(),
    }
    return render(request, 'matches/imprimir_reporte_tarjetas.html', context)


@login_required
def pagar_multas_equipo_torneo(request, torneo_id, equipo_id):
    """
    Permite cobrar todas las multas pendientes de un equipo para este torneo
    en un solo clic y registrarlas directamente en la Caja de Tesorería.
    """
    es_admin = bool(
        request.user.is_superuser or 
        getattr(request.user, 'role', '') in ['tesorero', 'tesoreria', 'superadmin', 'comision', 'organizador']
    )
    if not es_admin:
        messages.error(request, "No tienes autorización para cobrar multas en tesorería.")
        return redirect('control_tarjetas_torneo', torneo_id=torneo_id)

    torneo = _get_torneo_safe(request, torneo_id)
    org = getattr(request, 'organizacion', None) or torneo.organizacion
    equipo = get_object_or_404(Equipo, id=equipo_id)

    if request.method == 'POST':
        multas_filter = {
            'partido__torneo': torneo,
            'equipo': equipo,
            'estado': 'pendiente'
        }
        if org:
            multas_filter['organizacion'] = org

        multas_pendientes = MultaTarjeta.objects.filter(**multas_filter)
        total_monto = multas_pendientes.aggregate(models.Sum('monto'))['monto__sum'] or Decimal('0.00')
        count = multas_pendientes.count()

        if count == 0:
            messages.info(request, f"El club {equipo.nombre} no tiene multas pendientes de pago en este torneo.")
            return redirect('control_tarjetas_torneo', torneo_id=torneo.id)

        now = timezone.now()
        with transaction.atomic():
            multas_pendientes.update(estado='pagado', fecha_pago=now)

            if org:
                MovimientoCaja.objects.create(
                    organizacion=org,
                    tipo='ingreso',
                    monto=total_monto,
                    concepto=f"Cobro Tarjetas ({count} sanciones) - Club: {equipo.nombre} (Torneo: {torneo.nombre})",
                    registrado_por=request.user
                )

        messages.success(
            request, 
            f"¡Pago registrado exitosamente en Tesorería! Se cobraron {count} multa(s) del club {equipo.nombre} por un total de ${total_monto:.2f}."
        )

    return redirect('control_tarjetas_torneo', torneo_id=torneo.id)


@login_required
def exportar_excel_tarjetas_torneo(request, torneo_id):
    """
    Exporta a Excel (.xlsx) el reporte financiero y disciplinario de tarjetas del torneo:
    - Si tipo='equipo' y equipo_id: genera el Estado de Cuenta detallado del club.
    - Si tipo='jornada' y jornada: genera la Planilla Disciplinaria de esa fecha/programación.
    - Por defecto: genera el libro consolidado con 2 pestañas (Resumen por Equipo + Historial).
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    torneo = _get_torneo_safe(request, torneo_id)
    org = getattr(request, 'organizacion', None) or torneo.organizacion
    datos = _calcular_datos_tarjetas(torneo, org)
    multas_qs = datos['multas_qs']

    tipo = request.GET.get('tipo', 'total')
    equipo_id_raw = request.GET.get('equipo_id', '')
    jornada_raw = request.GET.get('jornada', '')

    wb = openpyxl.Workbook()

    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_header_gold = PatternFill(start_color="D97706", end_color="D97706", fill_type="solid")
    fill_header_green = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # 1. Exportación filtrada por Equipo
    if tipo == 'equipo' and equipo_id_raw.isdigit():
        eid = int(equipo_id_raw)
        ed = next((x for x in datos['equipos_data'] if x['equipo'] and x['equipo'].id == eid), None)
        if ed:
            ws = wb.active
            ws.title = f"Multas - {ed['equipo'].nombre[:25]}"
            ws.views.sheetView[0].showGridLines = True
            ws.append([f"ESTADO DE CUENTA DE SANCIONES Y MULTAS - {ed['equipo'].nombre.upper()}"])
            ws.cell(row=1, column=1).font = Font(name="Arial", size=13, bold=True, color="1E3A8A")
            ws.append([f"Torneo: {torneo.nombre} | Grupo: {ed['grupo_nombre']} | Emisión: {timezone.now().strftime('%d/%m/%Y %H:%M')}"])
            ws.append([f"Tarjetas: {ed['total_tarjetas']} ({ed['amarillas']} Amarillas, {ed['rojas']} Rojas) | Total: ${ed['monto_total']} | Pagado: ${ed['monto_pagado']} | Pendiente: ${ed['monto_pendiente']}"])
            ws.append([])

            headers = ["N° Multa", "Fecha / Jornada", "Partido / Rival", "Minuto", "Tarjeta", "Jugador Sancionado", "Monto ($)", "Estado", "Fecha de Pago"]
            ws.append(headers)
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=5, column=col_num)
                cell.font = font_header
                cell.fill = fill_header_green
                cell.alignment = Alignment(horizontal="center", vertical="center")

            row_idx = 6
            for m in ed['multas_list']:
                p = m.partido
                if p:
                    rival = p.equipo_visitante.nombre if p.equipo_local_id == ed['equipo'].id else p.equipo_local.nombre
                    p_txt = f"vs {rival}"
                    f_txt = f"Fecha {p.jornada}" if p.jornada else "-"
                else:
                    p_txt = "Sanción Administrativa"
                    f_txt = "-"
                min_txt = f"{m.evento.minuto}'" if m.evento and m.evento.minuto else "-"
                jug_txt = (m.jugador.get_full_name() or m.jugador.username) if m.jugador else "-"
                pago_txt = m.fecha_pago.strftime('%d/%m/%Y %H:%M') if m.fecha_pago else "-"

                ws.append([
                    m.id,
                    f_txt,
                    p_txt,
                    min_txt,
                    m.get_motivo_display(),
                    jug_txt,
                    float(m.monto),
                    m.get_estado_display(),
                    pago_txt
                ])
                for c in range(1, 10):
                    cell = ws.cell(row=row_idx, column=c)
                    cell.border = thin_border
                    if c in [1, 2, 4, 5, 8, 9]:
                        cell.alignment = Alignment(horizontal="center")
                    elif c == 7:
                        cell.number_format = "$#,##0.00"
                        cell.alignment = Alignment(horizontal="right")
                row_idx += 1

            for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']:
                ws.column_dimensions[col].width = 18
            ws.column_dimensions['C'].width = 25
            ws.column_dimensions['F'].width = 28

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            filename = f"Multas_{ed['equipo'].nombre.replace(' ', '_')}_{timezone.now().strftime('%Y%m%d')}.xlsx"
            response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            return response

    # 2. Exportación filtrada por Jornada / Programación
    if tipo == 'jornada' and jornada_raw.isdigit():
        jnum = int(jornada_raw)
        fd = next((x for x in datos['fechas_data'] if x['numero_fecha'] == jnum), None)
        if fd:
            ws = wb.active
            ws.title = f"Multas Fecha {jnum}"
            ws.views.sheetView[0].showGridLines = True
            ws.append([f"PLANILLA DE SANCIONES DISCIPLINARIAS - FECHA {jnum}"])
            ws.cell(row=1, column=1).font = Font(name="Arial", size=13, bold=True, color="1E3A8A")
            ws.append([f"Torneo: {torneo.nombre} | Emisión: {timezone.now().strftime('%d/%m/%Y %H:%M')}"])
            ws.append([f"Amarillas: {fd['total_amarillas']} | Rojas: {fd['total_rojas']} | Total Multas: ${fd['total_multas_monto']} | Pagado: ${fd['total_pagado_monto']} | Pendiente: ${fd['total_pendiente_monto']}"])
            ws.append([])

            headers = ["Partido", "Grupo", "Min", "Tarjeta", "Jugador", "Equipo", "Valor ($)", "Estado"]
            ws.append(headers)
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=5, column=col_num)
                cell.font = font_header
                cell.fill = fill_header_gold
                cell.alignment = Alignment(horizontal="center", vertical="center")

            row_idx = 6
            for p_item in fd['partidos']:
                p = p_item['partido']
                p_label = f"{p.equipo_local.nombre} vs {p.equipo_visitante.nombre}"
                g_label = p.grupo.nombre if p.grupo else "General"
                for ev_info in p_item['eventos']:
                    ev = ev_info['evento']
                    min_txt = f"{ev.minuto}'" if ev.minuto else "-"
                    jug_txt = (ev.jugador.get_full_name() or ev.jugador.username) if ev.jugador else "-"
                    eq_txt = ev.equipo.nombre if ev.equipo else "-"
                    monto_val = float(ev_info['monto'])
                    estado_txt = "Pagado" if ev_info['pagada'] else "Pendiente"

                    ws.append([
                        p_label,
                        g_label,
                        min_txt,
                        ev.get_tipo_display() if hasattr(ev, 'get_tipo_display') else str(ev.tipo).capitalize(),
                        jug_txt,
                        eq_txt,
                        monto_val,
                        estado_txt
                    ])
                    for c in range(1, 9):
                        cell = ws.cell(row=row_idx, column=c)
                        cell.border = thin_border
                        if c in [2, 3, 4, 8]:
                            cell.alignment = Alignment(horizontal="center")
                        elif c == 7:
                            cell.number_format = "$#,##0.00"
                            cell.alignment = Alignment(horizontal="right")
                    row_idx += 1

            for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H']:
                ws.column_dimensions[col].width = 20
            ws.column_dimensions['A'].width = 30
            ws.column_dimensions['E'].width = 25

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            filename = f"Multas_Fecha_{jnum}_{torneo.nombre.replace(' ', '_')}_{timezone.now().strftime('%Y%m%d')}.xlsx"
            response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            return response

    # 3. Exportación Consolidada General (2 Hojas)
    ws1 = wb.active
    ws1.title = "Resumen por Equipo"
    ws1.views.sheetView[0].showGridLines = True

    ws1.append([f"CONTROL DE TARJETAS Y MULTAS - {torneo.nombre.upper()}"])
    ws1.cell(row=1, column=1).font = Font(name="Arial", size=12, bold=True, color="1E3A8A")
    ws1.append([f"Amarilla: ${torneo.costo_amarilla} | Roja: ${torneo.costo_roja} | Generado: {timezone.now().strftime('%d/%m/%Y %H:%M')}"])
    ws1.append([])

    headers1 = ["Equipo", "Grupo", "Amarillas", "Rojas", "Total Tarjetas", "Total Multas ($)", "Cobrado ($)", "Pendiente ($)", "Estado"]
    ws1.append(headers1)
    for col_num in range(1, len(headers1) + 1):
        cell = ws1.cell(row=4, column=col_num)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_idx = 5
    for ed in datos['equipos_data']:
        estado_txt = "Al Día" if ed['monto_pendiente'] == Decimal('0.00') else "Pendiente"
        ws1.append([
            ed['equipo'].nombre if ed['equipo'] else "Equipo",
            ed['grupo_nombre'],
            ed['amarillas'],
            ed['rojas'],
            ed['total_tarjetas'],
            float(ed['monto_total']),
            float(ed['monto_pagado']),
            float(ed['monto_pendiente']),
            estado_txt
        ])
        for c in range(1, 10):
            cell = ws1.cell(row=row_idx, column=c)
            cell.border = thin_border
            if c in [3, 4, 5]:
                cell.alignment = Alignment(horizontal="center")
            elif c in [6, 7, 8]:
                cell.number_format = "$#,##0.00"
                cell.alignment = Alignment(horizontal="right")
        row_idx += 1

    for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']:
        ws1.column_dimensions[col].width = 16
    ws1.column_dimensions['A'].width = 28

    # HOJA 2: Historial Completo
    ws2 = wb.create_sheet(title="Historial de Amonestaciones")
    ws2.views.sheetView[0].showGridLines = True
    headers2 = ["ID", "Fecha / Jornada", "Partido", "Minuto", "Tarjeta", "Deportista Sancionado", "Equipo", "Monto ($)", "Estado", "Fecha de Pago"]
    ws2.append(headers2)
    for col_num in range(1, len(headers2) + 1):
        cell = ws2.cell(row=1, column=col_num)
        cell.font = font_header
        cell.fill = fill_header_gold
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for m in multas_qs.order_by('-partido__jornada', '-partido__fecha_hora', '-evento__minuto', '-id'):
        p = m.partido
        p_txt = f"{p.equipo_local.nombre} vs {p.equipo_visitante.nombre}" if p else "N/A"
        f_txt = f"Fecha {p.jornada}" if p and p.jornada else "-"
        min_txt = f"{m.evento.minuto}'" if m.evento and m.evento.minuto else "-"
        jug_txt = (m.jugador.get_full_name() or m.jugador.username) if m.jugador else "-"
        eq_txt = m.equipo.nombre if m.equipo else "-"
        pago_txt = m.fecha_pago.strftime('%d/%m/%Y %H:%M') if m.fecha_pago else "-"

        ws2.append([
            m.id,
            f_txt,
            p_txt,
            min_txt,
            m.get_motivo_display(),
            jug_txt,
            eq_txt,
            float(m.monto),
            m.get_estado_display(),
            pago_txt
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"Tarjetas_Multas_{torneo.nombre.replace(' ', '_')}_{timezone.now().strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
