import io
from django.db.models import Q, Count, Sum
from matches.models import Torneo, GrupoTorneo, Partido, EventoPartido, EquipoGrupoTorneo
from teams.models import Equipo, FichaJugador

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


def _resolver_desempate_grupos(equipos_stats, partidos_finalizados):
    """
    Aplica los criterios de desempate en orden:
    1. PTS (Puntos)
    2. DG (Diferencia de Goles)
    3. GF (Goles a Favor)
    4. Enfrentamiento directo (mini-tabla: Puntos en partidos entre los empatados)
    5. Diferencia de Goles en enfrentamiento directo
    6. Menor cantidad de tarjetas rojas
    7. Menor cantidad de tarjetas amarillas
    8. Empate pendiente de resolución administrativa (orden alfabético para estabilidad visual)
    """
    if not equipos_stats:
        return []

    # 1. Agrupar por (PTS, DG, GF)
    grupos_primer_nivel = {}
    for item in equipos_stats:
        key = (item['PTS'], item['DG'], item['GF'])
        grupos_primer_nivel.setdefault(key, []).append(item)

    claves_ordenadas = sorted(grupos_primer_nivel.keys(), key=lambda k: (k[0], k[1], k[2]), reverse=True)

    resultado_final = []

    for key in claves_ordenadas:
        subgrupo = grupos_primer_nivel[key]
        if len(subgrupo) == 1:
            resultado_final.append(subgrupo[0])
            continue

        # Subgrupo de 2 o más equipos empatados en PTS, DG y GF -> Mini-tabla de Enfrentamientos Directos
        ids_empatados = set(e['equipo_id'] for e in subgrupo)
        
        h2h_stats = {eid: {'h2h_pts': 0, 'h2h_dg': 0, 'h2h_gf': 0} for eid in ids_empatados}
        
        for p in partidos_finalizados:
            loc_id = p.equipo_local_id
            vis_id = p.equipo_visitante_id
            if loc_id in ids_empatados and vis_id in ids_empatados:
                gl = p.goles_local
                gv = p.goles_visitante
                
                h2h_stats[loc_id]['h2h_gf'] += gl
                h2h_stats[vis_id]['h2h_gf'] += gv
                h2h_stats[loc_id]['h2h_dg'] += (gl - gv)
                h2h_stats[vis_id]['h2h_dg'] += (gv - gl)
                
                if gl > gv:
                    h2h_stats[loc_id]['h2h_pts'] += 3
                elif gl < gv:
                    h2h_stats[vis_id]['h2h_pts'] += 3
                else:
                    h2h_stats[loc_id]['h2h_pts'] += 1
                    h2h_stats[vis_id]['h2h_pts'] += 1

        for item in subgrupo:
            eid = item['equipo_id']
            item['h2h_pts'] = h2h_stats[eid]['h2h_pts']
            item['h2h_dg'] = h2h_stats[eid]['h2h_dg']

        grupos_h2h = {}
        for item in subgrupo:
            h2h_key = (item['h2h_pts'], item['h2h_dg'])
            grupos_h2h.setdefault(h2h_key, []).append(item)

        claves_h2h = sorted(grupos_h2h.keys(), key=lambda k: (k[0], k[1]), reverse=True)

        for h_key in claves_h2h:
            sub_h2h = grupos_h2h[h_key]
            if len(sub_h2h) == 1:
                resultado_final.append(sub_h2h[0])
                continue

            sub_h2h_sorted = sorted(
                sub_h2h, 
                key=lambda x: (x['tarjetas_rojas'], x['tarjetas_amarillas'])
            )

            grupos_disciplina = {}
            for item in sub_h2h_sorted:
                disc_key = (item['tarjetas_rojas'], item['tarjetas_amarillas'])
                grupos_disciplina.setdefault(disc_key, []).append(item)

            for d_key, empatados_finales in grupos_disciplina.items():
                if len(empatados_finales) > 1:
                    for item in empatados_finales:
                        item['empate_pendiente'] = True
                    empatados_finales.sort(key=lambda x: x['nombre'])
                resultado_final.extend(empatados_finales)

    return resultado_final


def calcular_posiciones_grupo(grupo, torneo, organizacion):
    """
    Calcula la tabla de posiciones independiente de un grupo en un torneo personalizado.
    """
    if torneo.tipo != 'personalizado':
        raise ValueError("El cálculo de estadísticas personalizado solo aplica a torneos de tipo 'personalizado'.")

    asignaciones = EquipoGrupoTorneo.objects.filter(
        grupo=grupo,
        torneo=torneo,
        torneo__organizacion=organizacion
    ).select_related('equipo')
    
    equipos = [a.equipo for a in asignaciones]

    partidos_finalizados = list(
        Partido.objects.filter(
            organizacion=organizacion,
            torneo=torneo,
            grupo_personalizado=grupo,
            fase='grupos',
            estado='finalizado'
        ).select_related('equipo_local', 'equipo_visitante').order_by('fecha_hora', 'id')
    )

    pts_win = getattr(torneo, 'puntos_victoria', 3)
    pts_draw = getattr(torneo, 'puntos_empate', 1)
    pts_loss = getattr(torneo, 'puntos_derrota', 0)

    stats_map = {}
    for equipo in equipos:
        stats_map[equipo.id] = {
            'equipo': equipo,
            'equipo_id': equipo.id,
            'nombre': equipo.nombre,
            'logo': equipo.logo,
            'PJ': 0,
            'PG': 0,
            'PE': 0,
            'PP': 0,
            'GF': 0,
            'GC': 0,
            'DG': 0,
            'PTS': 0,
            'rendimiento': 0.0,
            'forma': [],
            'tarjetas_amarillas': 0,
            'tarjetas_rojas': 0,
            'posicion': 0,
            'es_clasificado_provisional': False,
            'empate_pendiente': False,
        }

    for p in partidos_finalizados:
        loc_id = p.equipo_local_id
        vis_id = p.equipo_visitante_id

        gl = p.goles_local
        gv = p.goles_visitante

        if loc_id in stats_map:
            s_loc = stats_map[loc_id]
            s_loc['PJ'] += 1
            s_loc['GF'] += gl
            s_loc['GC'] += gv
            s_loc['DG'] = s_loc['GF'] - s_loc['GC']

            if gl > gv:
                s_loc['PG'] += 1
                s_loc['PTS'] += pts_win
                s_loc['forma'].append('G')
            elif gl < gv:
                s_loc['PP'] += 1
                s_loc['PTS'] += pts_loss
                s_loc['forma'].append('P')
            else:
                s_loc['PE'] += 1
                s_loc['PTS'] += pts_draw
                s_loc['forma'].append('E')

        if vis_id in stats_map:
            s_vis = stats_map[vis_id]
            s_vis['PJ'] += 1
            s_vis['GF'] += gv
            s_vis['GC'] += gl
            s_vis['DG'] = s_vis['GF'] - s_vis['GC']

            if gv > gl:
                s_vis['PG'] += 1
                s_vis['PTS'] += pts_win
                s_vis['forma'].append('G')
            elif gv < gl:
                s_vis['PP'] += 1
                s_vis['PTS'] += pts_loss
                s_vis['forma'].append('P')
            else:
                s_vis['PE'] += 1
                s_vis['PTS'] += pts_draw
                s_vis['forma'].append('E')

    eventos_tarjetas = EventoPartido.objects.filter(
        partido__in=partidos_finalizados,
        tipo__in=['amarilla', 'roja']
    )

    for ev in eventos_tarjetas:
        if ev.equipo_id and ev.equipo_id in stats_map:
            if ev.tipo == 'amarilla':
                stats_map[ev.equipo_id]['tarjetas_amarillas'] += 1
            elif ev.tipo == 'roja':
                stats_map[ev.equipo_id]['tarjetas_rojas'] += 1

    for item in stats_map.values():
        pj = item['PJ']
        pts_max = pj * pts_win
        item['rendimiento'] = round((item['PTS'] / pts_max * 100), 1) if pts_max > 0 else 0.0
        item['forma'] = item['forma'][-5:]

    equipos_ordenados = _resolver_desempate_grupos(list(stats_map.values()), partidos_finalizados)

    cupos = grupo.cupos_clasificacion
    for idx, item in enumerate(equipos_ordenados, start=1):
        item['posicion'] = idx
        if idx <= cupos:
            item['es_clasificado_provisional'] = True

    advertencia_cupos = None
    if cupos > len(equipos) and len(equipos) > 0:
        advertencia_cupos = f"Este grupo posee {len(equipos)} equipos y tiene configurados {cupos} cupos de clasificación. Revisa la configuración."

    return {
        'grupo': grupo,
        'tabla': equipos_ordenados,
        'cupos': cupos,
        'advertencia_cupos': advertencia_cupos,
    }


def calcular_estadisticas_jugadores_grupo(grupo, torneo, organizacion):
    """
    Calcula goleadores, asistencias, tarjetas amarillas y rojas de un grupo específico.
    Filtra únicamente eventos de partidos finalizados del grupo y torneo.
    """
    partidos_finalizados = Partido.objects.filter(
        organizacion=organizacion,
        torneo=torneo,
        grupo_personalizado=grupo,
        fase='grupos',
        estado='finalizado'
    )

    eventos = EventoPartido.objects.filter(
        partido__in=partidos_finalizados,
        jugador__isnull=False
    ).select_related('jugador', 'equipo')

    fichas_map = {}
    fichas_qs = FichaJugador.objects.filter(
        organizacion=organizacion
    ).select_related('user', 'equipo')

    for f in fichas_qs:
        if f.torneo_id == torneo.id and f.equipo_id:
            fichas_map[(f.user_id, f.equipo_id, f.torneo_id)] = f
        if f.equipo_id and (f.user_id, f.equipo_id, None) not in fichas_map:
            fichas_map[(f.user_id, f.equipo_id, None)] = f
        if (f.user_id, None, None) not in fichas_map:
            fichas_map[(f.user_id, None, None)] = f

    def get_ficha_jugador(user_id, equipo_id):
        return (
            fichas_map.get((user_id, equipo_id, torneo.id)) or
            fichas_map.get((user_id, equipo_id, None)) or
            fichas_map.get((user_id, None, None))
        )

    conteo_goles = {}
    conteo_asistencias = {}
    conteo_amarillas = {}
    conteo_rojas = {}

    for ev in eventos:
        key = (ev.jugador_id, ev.equipo_id if ev.equipo_id else None)
        if ev.tipo == 'gol':
            conteo_goles[key] = conteo_goles.get(key, 0) + 1
        elif ev.tipo == 'asistencia':
            conteo_asistencias[key] = conteo_asistencias.get(key, 0) + 1
        elif ev.tipo == 'amarilla':
            conteo_amarillas[key] = conteo_amarillas.get(key, 0) + 1
        elif ev.tipo == 'roja':
            conteo_rojas[key] = conteo_rojas.get(key, 0) + 1

    def construir_ranking(conteo_dict):
        items = []
        for (user_id, equipo_id), total in conteo_dict.items():
            if total <= 0:
                continue
            ficha = get_ficha_jugador(user_id, equipo_id)
            nombre_jugador = ficha.user.get_full_name() if (ficha and ficha.user) else "Jugador N/A"
            equipo_obj = ficha.equipo if (ficha and ficha.equipo) else None
            
            if not equipo_obj and equipo_id:
                try:
                    equipo_obj = Equipo.objects.get(id=equipo_id)
                except Equipo.DoesNotExist:
                    equipo_obj = None

            foto_url = ficha.foto.url if (ficha and ficha.foto) else None
            dorsal = ficha.numero_camiseta if (ficha and ficha.numero_camiseta is not None) else "-"
            
            partidos_jugados = 0
            if equipo_obj:
                partidos_jugados = partidos_finalizados.filter(
                    Q(equipo_local=equipo_obj) | Q(equipo_visitante=equipo_obj)
                ).count()

            items.append({
                'user_id': user_id,
                'nombre_jugador': nombre_jugador,
                'foto_url': foto_url,
                'dorsal': dorsal,
                'equipo': equipo_obj,
                'equipo_nombre': equipo_obj.nombre if equipo_obj else "Sin Equipo",
                'equipo_logo': equipo_obj.logo if equipo_obj else None,
                'total': total,
                'partidos_jugados': partidos_jugados,
            })
        
        items.sort(key=lambda x: (-x['total'], x['nombre_jugador']))
        for idx, item in enumerate(items, start=1):
            item['posicion'] = idx
        return items

    return {
        'goleadores': construir_ranking(conteo_goles),
        'asistencias': construir_ranking(conteo_asistencias),
        'amarillas': construir_ranking(conteo_amarillas),
        'rojas': construir_ranking(conteo_rojas),
    }


def calcular_resumen_grupo(grupo, torneo, organizacion):
    """
    Calcula el resumen de avance y métricas generales del grupo.
    """
    equipos_count = EquipoGrupoTorneo.objects.filter(
        grupo=grupo, torneo=torneo
    ).count()

    partidos_qs = Partido.objects.filter(
        organizacion=organizacion,
        torneo=torneo,
        grupo_personalizado=grupo,
        fase='grupos'
    )

    partidos_totales = partidos_qs.count()
    partidos_finalizados_qs = partidos_qs.filter(estado='finalizado')
    partidos_finalizados = partidos_finalizados_qs.count()
    partidos_pendientes = partidos_totales - partidos_finalizados

    porcentaje_avance = round((partidos_finalizados / partidos_totales * 100), 1) if partidos_totales > 0 else 0.0

    if partidos_totales == 0:
        estado_grupo = 'Sin fixture'
    elif partidos_finalizados == 0:
        estado_grupo = 'Pendiente'
    elif partidos_finalizados < partidos_totales:
        estado_grupo = 'En competencia'
    else:
        estado_grupo = 'Finalizado'

    goles_totales = 0
    for p in partidos_finalizados_qs:
        goles_totales += (p.goles_local + p.goles_visitante)

    promedio_goles = round((goles_totales / partidos_finalizados), 2) if partidos_finalizados > 0 else 0.0

    eventos_tarjetas = EventoPartido.objects.filter(
        partido__in=partidos_finalizados_qs,
        tipo__in=['amarilla', 'roja']
    )
    amarillas_totales = eventos_tarjetas.filter(tipo='amarilla').count()
    rojas_totales = eventos_tarjetas.filter(tipo='roja').count()

    return {
        'equipos_count': equipos_count,
        'partidos_totales': partidos_totales,
        'partidos_finalizados': partidos_finalizados,
        'partidos_pendientes': partidos_pendientes,
        'porcentaje_avance': porcentaje_avance,
        'estado_grupo': estado_grupo,
        'goles_totales': goles_totales,
        'promedio_goles': promedio_goles,
        'amarillas_totales': amarillas_totales,
        'rojas_totales': rojas_totales,
        'cupos_clasificacion': grupo.cupos_clasificacion,
    }


def obtener_partidos_recientes_y_proximos(grupo, torneo, organizacion):
    """
    Obtiene los partidos recientemente finalizados y los próximos partidos programados del grupo.
    """
    partidos_base = Partido.objects.filter(
        organizacion=organizacion,
        torneo=torneo,
        grupo_personalizado=grupo,
        fase='grupos'
    ).select_related('equipo_local', 'equipo_visitante', 'arbitro', 'vocal')

    recientes = partidos_base.filter(
        estado='finalizado'
    ).order_by('-fecha_hora', '-id')[:5]

    proximos = partidos_base.filter(
        estado__in=['programado', 'en_curso']
    ).order_by('fecha_hora', 'id')[:5]

    return {
        'recientes me': recientes, # Key alias standard
        'recientes': recientes,
        'proximos': proximos,
    }


def calcular_resumen_torneo_general(torneo, organizacion):
    """
    Calcula la vista general combinada del torneo de grupos.
    """
    grupos = list(GrupoTorneo.objects.filter(torneo=torneo).order_by('orden', 'id'))
    
    total_equipos = EquipoGrupoTorneo.objects.filter(torneo=torneo).count()
    partidos_qs = Partido.objects.filter(organizacion=organizacion, torneo=torneo, fase='grupos')
    total_partidos = partidos_qs.count()
    partidos_finalizados = partidos_qs.filter(estado='finalizado').count()
    partidos_pendientes = total_partidos - partidos_finalizados
    avance_general = round((partidos_finalizados / total_partidos * 100), 1) if total_partidos > 0 else 0.0

    goles_totales = 0
    for p in partidos_qs.filter(estado='finalizado'):
        goles_totales += (p.goles_local + p.goles_visitante)

    tarjetas_qs = EventoPartido.objects.filter(partido__in=partidos_qs.filter(estado='finalizado'), tipo__in=['amarilla', 'roja'])
    total_amarillas = tarjetas_qs.filter(tipo='amarilla').count()
    total_rojas = tarjetas_qs.filter(tipo='roja').count()

    resumen_grupos = []
    for g in grupos:
        res = calcular_posiciones_grupo(g, torneo, organizacion)
        res_info = calcular_resumen_grupo(g, torneo, organizacion)
        
        lider = res['tabla'][0] if res['tabla'] else None
        
        resumen_grupos.append({
            'grupo': g,
            'resumen': res_info,
            'lider': lider,
            'clasificados': [item for item in res['tabla'] if item['es_clasificado_provisional']],
            'tabla': res['tabla'],
        })

    return {
        'torneo': torneo,
        'total_grupos': len(grupos),
        'total_equipos': total_equipos,
        'total_partidos': total_partidos,
        'partidos_finalizados': partidos_finalizados,
        'partidos_pendientes': partidos_pendientes,
        'avance_general': avance_general,
        'goles_totales': goles_totales,
        'total_amarillas': total_amarillas,
        'total_rojas': total_rojas,
        'resumen_grupos': resumen_grupos,
    }


def generar_excel_estadisticas_torneo(torneo, organizacion, grupo_id=None):
    """
    Genera un libro de Excel en memoria (.xlsx) usando openpyxl con la información del torneo.
    """
    if not OPENPYXL_AVAILABLE:
        raise ImportError("La librería openpyxl no está instalada.")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="4B5563")
    bold_font = Font(name="Calibri", size=11, bold=True)
    normal_font = Font(name="Calibri", size=11)
    clasificado_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    grupos_qs = GrupoTorneo.objects.filter(torneo=torneo).order_by('orden', 'id')
    if grupo_id:
        grupos_qs = grupos_qs.filter(id=grupo_id)

    grupos = list(grupos_qs)

    # 1. Hoja de Resumen General
    ws_resumen = wb.create_sheet(title="Resumen")
    ws_resumen.views.sheetView[0].showGridLines = True
    
    ws_resumen.append([f"RESUMEN DE POSICIONES Y ESTADÍSTICAS - {torneo.nombre.upper()}"])
    ws_resumen.append([f"Organización: {organizacion.nombre} | Categoría: {torneo.get_categoria_display()} | Temporada: {torneo.temporada}"])
    ws_resumen.append([])

    ws_resumen['A1'].font = title_font
    ws_resumen['A2'].font = subtitle_font

    general_data = calcular_resumen_torneo_general(torneo, organizacion)

    ws_resumen.append(["Métrica", "Valor"])
    ws_resumen.append(["Total de Grupos", general_data['total_grupos']])
    ws_resumen.append(["Total de Equipos", general_data['total_equipos']])
    ws_resumen.append(["Partidos Totales", general_data['total_partidos']])
    ws_resumen.append(["Partidos Finalizados", general_data['partidos_finalizados']])
    ws_resumen.append(["Partidos Pendientes", general_data['partidos_pendientes']])
    ws_resumen.append(["Porcentaje de Avance", f"{general_data['avance_general']}%"])
    ws_resumen.append(["Total Goles Anotados", general_data['goles_totales']])
    ws_resumen.append(["Total Tarjetas Amarillas", general_data['total_amarillas']])
    ws_resumen.append(["Total Tarjetas Rojas", general_data['total_rojas']])

    for r in range(4, 14):
        ws_resumen.cell(row=r, column=1).font = bold_font
        ws_resumen.cell(row=r, column=1).border = thin_border
        ws_resumen.cell(row=r, column=2).font = normal_font
        ws_resumen.cell(row=r, column=2).border = thin_border

    ws_resumen.cell(row=4, column=1).fill = header_fill
    ws_resumen.cell(row=4, column=1).font = header_font
    ws_resumen.cell(row=4, column=2).fill = header_fill
    ws_resumen.cell(row=4, column=2).font = header_font

    # 2. Hojas por Grupo
    for g in grupos:
        sheet_name = f"Grupo {g.nombre}"[:31]
        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True

        ws.append([f"TABLA DE POSICIONES - {g.nombre} ({g.sector or 'Sin Sector'})"])
        ws.append([f"Torneo: {torneo.nombre} | Cupos Clasificación: {g.cupos_clasificacion}"])
        ws.append([])

        ws['A1'].font = title_font
        ws['A2'].font = subtitle_font

        headers = ["POS", "EQUIPO", "PJ", "PG", "PE", "PP", "GF", "GC", "DG", "PTS", "REND %", "CLASIFICACIÓN"]
        ws.append(headers)

        header_row = 4
        for col_num, h_text in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        pos_data = calcular_posiciones_grupo(g, torneo, organizacion)

        for item in pos_data['tabla']:
            clas_text = "Clasificado Provisional" if item['es_clasificado_provisional'] else "Fuera de Clasificación"
            if item['empate_pendiente']:
                clas_text += " (Empate Pendiente)"

            row_values = [
                item['posicion'],
                item['nombre'],
                item['PJ'],
                item['PG'],
                item['PE'],
                item['PP'],
                item['GF'],
                item['GC'],
                item['DG'],
                item['PTS'],
                f"{item['rendimiento']}%",
                clas_text
            ]
            ws.append(row_values)

            current_row = ws.max_row
            is_clas = item['es_clasificado_provisional']

            for c_idx in range(1, len(row_values) + 1):
                cell = ws.cell(row=current_row, column=c_idx)
                cell.font = normal_font
                cell.border = thin_border
                if is_clas:
                    cell.fill = clasificado_fill
                if c_idx not in (2, 12):
                    cell.alignment = Alignment(horizontal="center")

        ws.freeze_panes = "A5"

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # 3. Hojas de Líderes Individuales
    ws_gol = wb.create_sheet(title="Goleadores")
    ws_gol.views.sheetView[0].showGridLines = True
    ws_gol.append([f"TABLA DE GOLEADORES - {torneo.nombre.upper()}"])
    ws_gol.append(["POS", "JUGADOR", "DORSAL", "EQUIPO", "GRUPO", "GOLES TOTALES", "PARTIDOS JUGADOS"])
    ws_gol['A1'].font = title_font

    for c in range(1, 8):
        cell = ws_gol.cell(row=2, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for g in grupos:
        jug_data = calcular_estadisticas_jugadores_grupo(g, torneo, organizacion)
        for gol in jug_data['goleadores']:
            ws_gol.append([
                gol['posicion'],
                gol['nombre_jugador'],
                gol['dorsal'],
                gol['equipo_nombre'],
                g.nombre,
                gol['total'],
                gol['partidos_jugados']
            ])
            for c in range(1, 8):
                ws_gol.cell(row=ws_gol.max_row, column=c).border = thin_border

    # Hoja de Asistencias
    ws_asis = wb.create_sheet(title="Asistencias")
    ws_asis.views.sheetView[0].showGridLines = True
    ws_asis.append([f"TABLA DE ASISTENCIAS - {torneo.nombre.upper()}"])
    ws_asis.append(["POS", "JUGADOR", "DORSAL", "EQUIPO", "GRUPO", "ASISTENCIAS TOTALES", "PARTIDOS JUGADOS"])
    ws_asis['A1'].font = title_font

    for c in range(1, 8):
        cell = ws_asis.cell(row=2, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for g in grupos:
        jug_data = calcular_estadisticas_jugadores_grupo(g, torneo, organizacion)
        for asis in jug_data['asistencias']:
            ws_asis.append([
                asis['posicion'],
                asis['nombre_jugador'],
                asis['dorsal'],
                asis['equipo_nombre'],
                g.nombre,
                asis['total'],
                asis['partidos_jugados']
            ])
            for c in range(1, 8):
                ws_asis.cell(row=ws_asis.max_row, column=c).border = thin_border

    # Hoja de Disciplina
    ws_disc = wb.create_sheet(title="Disciplina")
    ws_disc.views.sheetView[0].showGridLines = True
    ws_disc.append([f"TABLA DE DISCIPLINA (TARJETAS) - {torneo.nombre.upper()}"])
    ws_disc.append(["POS", "JUGADOR", "DORSAL", "EQUIPO", "GRUPO", "AMARILLAS", "ROJAS"])
    ws_disc['A1'].font = title_font

    for c in range(1, 8):
        cell = ws_disc.cell(row=2, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for g in grupos:
        jug_data = calcular_estadisticas_jugadores_grupo(g, torneo, organizacion)
        # Combinar amarillas y rojas
        disc_map = {}
        for am in jug_data['amarillas']:
            k = (am['user_id'], am['nombre_jugador'], am['dorsal'], am['equipo_nombre'])
            disc_map[k] = {'amarillas': am['total'], 'rojas': 0}
        for rj in jug_data['rojas']:
            k = (rj['user_id'], rj['nombre_jugador'], rj['dorsal'], rj['equipo_nombre'])
            if k not in disc_map:
                disc_map[k] = {'amarillas': 0, 'rojas': rj['total']}
            else:
                disc_map[k]['rojas'] = rj['total']

        for idx, (k, val) in enumerate(sorted(disc_map.items(), key=lambda x: (-x[1]['rojas'], -x[1]['amarillas'])), 1):
            ws_disc.append([
                idx,
                k[1],
                k[2],
                k[3],
                g.nombre,
                val['amarillas'],
                val['rojas']
            ])
            for c in range(1, 8):
                ws_disc.cell(row=ws_disc.max_row, column=c).border = thin_border

    for ws_curr in [ws_resumen, ws_gol, ws_asis, ws_disc]:
        for col in ws_curr.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_curr.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def calcular_posiciones_general(torneo, organizacion):
    """
    Calcula la tabla de posiciones general para un torneo de tipo 'liga' o cualquier torneo sin grupos asignados.
    """
    equipos = list(torneo.equipos.all())
    partidos_all = list(
        Partido.objects.filter(
            organizacion=organizacion,
            torneo=torneo,
        ).select_related('equipo_local', 'equipo_visitante').order_by('fecha_hora', 'id')
    )
    
    equipos_map = {e.id: e for e in equipos}
    for p in partidos_all:
        if p.equipo_local and p.equipo_local_id not in equipos_map:
            equipos_map[p.equipo_local_id] = p.equipo_local
        if p.equipo_visitante and p.equipo_visitante_id not in equipos_map:
            equipos_map[p.equipo_visitante_id] = p.equipo_visitante

    partidos_finalizados = [p for p in partidos_all if p.estado == 'finalizado']

    pts_win = getattr(torneo, 'puntos_victoria', 3)
    pts_draw = getattr(torneo, 'puntos_empate', 1)
    pts_loss = getattr(torneo, 'puntos_derrota', 0)

    stats_map = {}
    for eq_id, equipo in equipos_map.items():
        stats_map[eq_id] = {
            'equipo': equipo,
            'equipo_id': equipo.id,
            'nombre': equipo.nombre,
            'logo': equipo.logo,
            'PJ': 0, 'PG': 0, 'PE': 0, 'PP': 0,
            'GF': 0, 'GC': 0, 'DG': 0, 'PTS': 0,
            'rendimiento': 0.0,
            'forma': [],
            'tarjetas_amarillas': 0,
            'tarjetas_rojas': 0,
            'posicion': 0,
            'es_clasificado_provisional': False,
            'empate_pendiente': False,
        }

    for p in partidos_finalizados:
        loc_id = p.equipo_local_id
        vis_id = p.equipo_visitante_id
        gl = p.goles_local
        gv = p.goles_visitante

        if loc_id in stats_map:
            s_loc = stats_map[loc_id]
            s_loc['PJ'] += 1
            s_loc['GF'] += gl
            s_loc['GC'] += gv
            s_loc['DG'] = s_loc['GF'] - s_loc['GC']
            if gl > gv:
                s_loc['PG'] += 1
                s_loc['PTS'] += pts_win
                s_loc['forma'].append('G')
            elif gl < gv:
                s_loc['PP'] += 1
                s_loc['PTS'] += pts_loss
                s_loc['forma'].append('P')
            else:
                s_loc['PE'] += 1
                s_loc['PTS'] += pts_draw
                s_loc['forma'].append('E')

        if vis_id in stats_map:
            s_vis = stats_map[vis_id]
            s_vis['PJ'] += 1
            s_vis['GF'] += gv
            s_vis['GC'] += gl
            s_vis['DG'] = s_vis['GF'] - s_vis['GC']
            if gv > gl:
                s_vis['PG'] += 1
                s_vis['PTS'] += pts_win
                s_vis['forma'].append('G')
            elif gv < gl:
                s_vis['PP'] += 1
                s_vis['PTS'] += pts_loss
                s_vis['forma'].append('P')
            else:
                s_vis['PE'] += 1
                s_vis['PTS'] += pts_draw
                s_vis['forma'].append('E')

    equipos_ordenados = _resolver_desempate_grupos(list(stats_map.values()), partidos_finalizados)

    for idx, item in enumerate(equipos_ordenados, start=1):
        item['posicion'] = idx

    return {
        'tabla': equipos_ordenados,
    }


def calcular_estadisticas_jugadores_general(torneo, organizacion):
    """
    Calcula goleadores, asistencias, tarjetas para un torneo general/liga.
    """
    partidos_finalizados = Partido.objects.filter(
        organizacion=organizacion,
        torneo=torneo,
        estado='finalizado'
    )
    eventos = EventoPartido.objects.filter(
        partido__in=partidos_finalizados,
        jugador__isnull=False
    ).select_related('jugador', 'equipo')

    fichas_map = {}
    fichas_qs = FichaJugador.objects.filter(
        organizacion=organizacion
    ).select_related('user', 'equipo')

    for f in fichas_qs:
        if f.torneo_id == torneo.id and f.equipo_id:
            fichas_map[(f.user_id, f.equipo_id, f.torneo_id)] = f
        if f.equipo_id and (f.user_id, f.equipo_id, None) not in fichas_map:
            fichas_map[(f.user_id, f.equipo_id, None)] = f
        if (f.user_id, None, None) not in fichas_map:
            fichas_map[(f.user_id, None, None)] = f

    def get_ficha_jugador(user_id, equipo_id):
        return (
            fichas_map.get((user_id, equipo_id, torneo.id)) or
            fichas_map.get((user_id, equipo_id, None)) or
            fichas_map.get((user_id, None, None))
        )

    conteo_goles = {}
    conteo_asistencias = {}
    conteo_amarillas = {}
    conteo_rojas = {}

    for ev in eventos:
        key = (ev.jugador_id, ev.equipo_id if ev.equipo_id else None)
        if ev.tipo == 'gol':
            conteo_goles[key] = conteo_goles.get(key, 0) + 1
        elif ev.tipo == 'asistencia':
            conteo_asistencias[key] = conteo_asistencias.get(key, 0) + 1
        elif ev.tipo == 'amarilla':
            conteo_amarillas[key] = conteo_amarillas.get(key, 0) + 1
        elif ev.tipo == 'roja':
            conteo_rojas[key] = conteo_rojas.get(key, 0) + 1

    def construir_ranking(conteo_dict):
        items = []
        for (user_id, equipo_id), total in conteo_dict.items():
            if total <= 0:
                continue
            ficha = get_ficha_jugador(user_id, equipo_id)
            nombre_jugador = ficha.user.get_full_name() if (ficha and ficha.user) else "Jugador N/A"
            equipo_obj = ficha.equipo if (ficha and ficha.equipo) else None
            
            if not equipo_obj and equipo_id:
                try:
                    equipo_obj = Equipo.objects.get(id=equipo_id)
                except Equipo.DoesNotExist:
                    equipo_obj = None

            foto_url = ficha.foto.url if (ficha and ficha.foto) else None
            dorsal = ficha.numero_camiseta if (ficha and ficha.numero_camiseta is not None) else "-"
            
            partidos_jugados = 0
            if equipo_obj:
                partidos_jugados = partidos_finalizados.filter(
                    Q(equipo_local=equipo_obj) | Q(equipo_visitante=equipo_obj)
                ).count()

            items.append({
                'user_id': user_id,
                'nombre_jugador': nombre_jugador,
                'foto_url': foto_url,
                'dorsal': dorsal,
                'equipo': equipo_obj,
                'equipo_nombre': equipo_obj.nombre if equipo_obj else "Sin Equipo",
                'equipo_logo': equipo_obj.logo if equipo_obj else None,
                'total': total,
                'partidos_jugados': partidos_jugados,
            })
        
        items.sort(key=lambda x: (-x['total'], x['nombre_jugador']))
        for idx, item in enumerate(items, start=1):
            item['posicion'] = idx
        return items

    return {
        'goleadores': construir_ranking(conteo_goles),
        'asistencias': construir_ranking(conteo_asistencias),
        'amarillas': construir_ranking(conteo_amarillas),
        'rojas': construir_ranking(conteo_rojas),
    }

