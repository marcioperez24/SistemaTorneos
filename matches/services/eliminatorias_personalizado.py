import io
from datetime import datetime, timedelta
from django.db import transaction
from django.core.exceptions import ValidationError
from django.utils import timezone
from django.utils.text import slugify

from matches.models import (
    Torneo, GrupoTorneo, Partido, EquipoGrupoTorneo,
    ClasificadoTorneo, ResolucionEmpateTorneo, LlaveEliminatoria,
    ResultadoFinalTorneo, BitacoraTorneo
)
from teams.models import Equipo

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


def evaluar_estado_llave(llave, organizacion):
    """
    Evalúa los partidos disputados en una llave eliminatoria y determina el resultado global,
    si existe un ganador deportivo o si se requiere penales / decisión administrativa.
    """
    if llave.torneo.tipo != 'personalizado':
        raise ValidationError("Esta función solo aplica a torneos personalizados.")

    if llave.es_bye:
        ganador = llave.equipo_local or llave.equipo_visitante
        return {
            'completa': True,
            'estado_sugerido': 'lista_para_confirmar',
            'ganador': ganador,
            'metodo': 'bye',
            'marcador_global_local': 0,
            'marcador_global_visitante': 0,
            'goles_visitante_local': 0,
            'goles_visitante_visitante': 0,
            'requiere_penales': False,
            'requiere_prorroga': False,
            'partidos': [],
            'mensaje': f"Llave con Pase Directo (BYE) para {ganador.nombre if ganador else 'Equipo'}."
        }

    if not llave.equipo_local or not llave.equipo_visitante:
        return {
            'completa': False,
            'estado_sugerido': 'pendiente',
            'ganador': None,
            'metodo': None,
            'requiere_penales': False,
            'partidos': [],
            'mensaje': "Esperando definición de ambos equipos."
        }

    # Buscar partidos de esta llave (filtrados por equipos y fase)
    partidos = list(
        Partido.objects.filter(
            organizacion=organizacion,
            torneo=llave.torneo,
            fase=llave.fase,
            equipo_local__in=[llave.equipo_local, llave.equipo_visitante],
            equipo_visitante__in=[llave.equipo_local, llave.equipo_visitante]
        ).order_by('numero_vuelta', 'fecha_hora', 'id')
    )

    partidos_finalizados = [p for p in partidos if p.estado == 'finalizado']

    if llave.formato == 'partido_unico':
        if len(partidos_finalizados) < 1:
            return {
                'completa': False,
                'estado_sugerido': 'programada' if partidos else 'pendiente',
                'ganador': None,
                'metodo': None,
                'requiere_penales': False,
                'partidos': partidos,
                'mensaje': "El partido único no ha finalizado."
            }

        p1 = partidos_finalizados[0]
        # Identificar goles según equipo_local / visitante de la llave
        if p1.equipo_local == llave.equipo_local:
            g_loc = p1.goles_local
            g_vis = p1.goles_visitante
        else:
            g_loc = p1.goles_visitante
            g_vis = p1.goles_local

        # Si ya se registraron penales en la llave
        if llave.definido_por_penales:
            if llave.penales_local is not None and llave.penales_visitante is not None:
                if llave.penales_local > llave.penales_visitante:
                    ganador = llave.equipo_local
                elif llave.penales_visitante > llave.penales_local:
                    ganador = llave.equipo_visitante
                else:
                    ganador = None

                return {
                    'completa': True,
                    'estado_sugerido': 'lista_para_confirmar',
                    'ganador': ganador,
                    'metodo': 'penales',
                    'marcador_global_local': g_loc,
                    'marcador_global_visitante': g_vis,
                    'goles_visitante_local': 0,
                    'goles_visitante_visitante': 0,
                    'definido_por_penales': True,
                    'penales_local': llave.penales_local,
                    'penales_visitante': llave.penales_visitante,
                    'hubo_prorroga': llave.hubo_prorroga,
                    'requiere_penales': False,
                    'partidos': partidos,
                    'mensaje': f"Definición por Penales: {llave.equipo_local.nombre} {llave.penales_local} - {llave.penales_visitante} {llave.equipo_visitante.nombre}"
                }

        if g_loc > g_vis:
            ganador = llave.equipo_local
            metodo = 'prorroga' if llave.hubo_prorroga else 'marcador'
            req_penales = False
        elif g_vis > g_loc:
            ganador = llave.equipo_visitante
            metodo = 'prorroga' if llave.hubo_prorroga else 'marcador'
            req_penales = False
        else:
            ganador = None
            metodo = None
            req_penales = True

        return {
            'completa': True if ganador else False,
            'estado_sugerido': 'lista_para_confirmar' if ganador else 'en_curso',
            'ganador': ganador,
            'metodo': metodo,
            'marcador_global_local': g_loc,
            'marcador_global_visitante': g_vis,
            'goles_visitante_local': 0,
            'goles_visitante_visitante': 0,
            'definido_por_penales': False,
            'hubo_prorroga': llave.hubo_prorroga,
            'requiere_penales': req_penales,
            'partidos': partidos,
            'mensaje': "Partido único finalizado." if ganador else "Empate registrado. Se requieren Penales o Prórroga."
        }

    else:
        # Ida y Vuelta
        if len(partidos_finalizados) < 2:
            return {
                'completa': False,
                'estado_sugerido': 'en_curso' if len(partidos_finalizados) == 1 else 'programada',
                'ganador': None,
                'metodo': None,
                'requiere_penales': False,
                'partidos': partidos,
                'mensaje': f"Se han disputado {len(partidos_finalizados)} de 2 partidos."
            }

        # Separar partido de ida (vuelta=1) y vuelta (vuelta=2)
        p_ida = next((p for p in partidos_finalizados if p.numero_vuelta == 1), partidos_finalizados[0])
        p_vta = next((p for p in partidos_finalizados if p.numero_vuelta == 2), partidos_finalizados[1])

        # Calcular goles globales por ID de equipo
        tot_goles = {llave.equipo_local.id: 0, llave.equipo_visitante.id: 0}
        goles_visitante = {llave.equipo_local.id: 0, llave.equipo_visitante.id: 0}

        # Partido Ida
        tot_goles[p_ida.equipo_local.id] += p_ida.goles_local
        tot_goles[p_ida.equipo_visitante.id] += p_ida.goles_visitante
        goles_visitante[p_ida.equipo_visitante.id] += p_ida.goles_visitante

        # Partido Vuelta
        tot_goles[p_vta.equipo_local.id] += p_vta.goles_local
        tot_goles[p_vta.equipo_visitante.id] += p_vta.goles_visitante
        goles_visitante[p_vta.equipo_visitante.id] += p_vta.goles_visitante

        gl_loc = tot_goles[llave.equipo_local.id]
        gl_vis = tot_goles[llave.equipo_visitante.id]
        gv_loc = goles_visitante[llave.equipo_local.id]
        gv_vis = goles_visitante[llave.equipo_visitante.id]

        if llave.definido_por_penales:
            if llave.penales_local is not None and llave.penales_visitante is not None:
                if llave.penales_local > llave.penales_visitante:
                    ganador = llave.equipo_local
                elif llave.penales_visitante > llave.penales_local:
                    ganador = llave.equipo_visitante
                else:
                    ganador = None

                return {
                    'completa': True,
                    'estado_sugerido': 'lista_para_confirmar',
                    'ganador': ganador,
                    'metodo': 'penales',
                    'marcador_global_local': gl_loc,
                    'marcador_global_visitante': gl_vis,
                    'goles_visitante_local': gv_loc,
                    'goles_visitante_visitante': gv_vis,
                    'definido_por_penales': True,
                    'penales_local': llave.penales_local,
                    'penales_visitante': llave.penales_visitante,
                    'hubo_prorroga': llave.hubo_prorroga,
                    'requiere_penales': False,
                    'partidos': partidos,
                    'mensaje': f"Definición por Penales: Global {gl_loc}-{gl_vis}, Penales {llave.penales_local}-{llave.penales_visitante}"
                }

        if gl_loc > gl_vis:
            ganador = llave.equipo_local
            metodo = 'marcador_global'
            req_penales = False
        elif gl_vis > gl_loc:
            ganador = llave.equipo_visitante
            metodo = 'marcador_global'
            req_penales = False
        else:
            # Empate global
            if llave.torneo.usar_gol_visitante:
                if gv_loc > gv_vis:
                    ganador = llave.equipo_local
                    metodo = 'gol_visitante'
                    req_penales = False
                elif gv_vis > gv_loc:
                    ganador = llave.equipo_visitante
                    metodo = 'gol_visitante'
                    req_penales = False
                else:
                    ganador = None
                    metodo = None
                    req_penales = True
            else:
                ganador = None
                metodo = None
                req_penales = True

        return {
            'completa': True if ganador else False,
            'estado_sugerido': 'lista_para_confirmar' if ganador else 'en_curso',
            'ganador': ganador,
            'metodo': metodo,
            'marcador_global_local': gl_loc,
            'marcador_global_visitante': gl_vis,
            'goles_visitante_local': gv_loc,
            'goles_visitante_visitante': gv_vis,
            'definido_por_penales': False,
            'hubo_prorroga': llave.hubo_prorroga,
            'requiere_penales': req_penales,
            'partidos': partidos,
            'mensaje': f"Global finalizado: {llave.equipo_local.nombre} {gl_loc} - {gl_vis} {llave.equipo_visitante.nombre}" if ganador else "Marcador global empatado. Se requieren Penales."
        }


def confirmar_ganador_llave(llave, ganador_id, usuario, organizacion, datos_confirmacion=None):
    """
    Confirma deportivamente o administrativamente el ganador de una llave eliminatoria.
    Actualiza la llave a 'finalizada', hace avanzar al ganador a la siguiente ronda
    y gestiona la creación del Tercer Lugar y Resultado Final (Campeón).
    """
    if llave.torneo.tipo != 'personalizado':
        raise ValidationError("Esta función solo aplica a torneos personalizados.")

    datos_confirmacion = datos_confirmacion or {}

    with transaction.atomic():
        llave_db = LlaveEliminatoria.objects.select_for_update().get(id=llave.id, organizacion=organizacion)

        if llave_db.estado == 'finalizada':
            raise ValidationError("Esta llave ya fue confirmada.")

        equipo_ganador = Equipo.objects.get(id=ganador_id, organizacion=organizacion)

        if equipo_ganador not in [llave_db.equipo_local, llave_db.equipo_visitante]:
            raise ValidationError(f"El equipo '{equipo_ganador.nombre}' no pertenece a esta llave.")

        # Si se envían penales manualmente en la confirmación, se aplican a la llave antes de evaluar
        if datos_confirmacion.get('penales_local') is not None and datos_confirmacion.get('penales_visitante') is not None:
            p_loc = int(datos_confirmacion['penales_local'])
            p_vis = int(datos_confirmacion['penales_visitante'])
            if p_loc < 0 or p_vis < 0:
                raise ValidationError("Los penales no pueden ser números negativos.")
            if p_loc == p_vis:
                raise ValidationError("La definición por penales no puede terminar en empate.")

            llave_db.definido_por_penales = True
            llave_db.penales_local = p_loc
            llave_db.penales_visitante = p_vis

        if datos_confirmacion.get('hubo_prorroga'):
            llave_db.hubo_prorroga = True

        eval_info = evaluar_estado_llave(llave_db, organizacion)

        motivo_admin = datos_confirmacion.get('motivo_admin')
        es_decision_admin = bool(motivo_admin)

        if not eval_info['completa'] and not llave_db.es_bye and not es_decision_admin:
            raise ValidationError(f"No se puede confirmar la llave: {eval_info['mensaje']}")

        if llave_db.definido_por_penales:
            p_loc = llave_db.penales_local
            p_vis = llave_db.penales_visitante
            ganador_penales = llave_db.equipo_local if p_loc > p_vis else llave_db.equipo_visitante

            if not es_decision_admin and equipo_ganador != ganador_penales:
                raise ValidationError(f"El ganador enviado ({equipo_ganador.nombre}) no coincide con el resultado de penales ({ganador_penales.nombre}).")

        elif not es_decision_admin and not llave_db.es_bye:
            if eval_info['ganador'] and equipo_ganador != eval_info['ganador']:
                raise ValidationError(f"El ganador enviado ({equipo_ganador.nombre}) no coincide con el ganador calculado ({eval_info['ganador'].nombre}).")

        # Asignar atributos a la llave
        llave_db.ganador = equipo_ganador
        llave_db.estado = 'finalizada'
        llave_db.confirmado_por = usuario
        llave_db.fecha_confirmacion = timezone.now()

        if es_decision_admin:
            llave_db.metodo_definicion = 'decision_administrativa'
        elif llave_db.es_bye:
            llave_db.metodo_definicion = 'bye'
        elif llave_db.definido_por_penales:
            llave_db.metodo_definicion = 'penales'
        else:
            llave_db.metodo_definicion = eval_info.get('metodo', 'marcador')

        llave_db.marcador_global_local = eval_info.get('marcador_global_local', 0)
        llave_db.marcador_global_visitante = eval_info.get('marcador_global_visitante', 0)
        llave_db.goles_visitante_local = eval_info.get('goles_visitante_local', 0)
        llave_db.goles_visitante_visitante = eval_info.get('goles_visitante_visitante', 0)
        llave_db.hubo_prorroga = datos_confirmacion.get('hubo_prorroga', eval_info.get('hubo_prorroga', False))

        llave_db.save()

        # Determine equipo perdedor
        equipo_perdedor = None
        if llave_db.equipo_local and llave_db.equipo_visitante:
            equipo_perdedor = llave_db.equipo_visitante if equipo_ganador == llave_db.equipo_local else llave_db.equipo_local

        # AVANCE A LA SIGUIENTE LLAVE
        if llave_db.siguiente_llave:
            sig = llave_db.siguiente_llave
            if llave_db.posicion_siguiente_llave == 'local':
                sig.equipo_local = equipo_ganador
            else:
                sig.equipo_visitante = equipo_ganador

            if sig.equipo_local and sig.equipo_visitante:
                sig.estado = 'programada'
            sig.save()

        # MANEJO DE TERCER LUGAR (Si se finaliza una semifinal y está activo disputar_tercer_lugar)
        if llave_db.fase == 'semifinal' and llave_db.torneo.disputar_tercer_lugar and equipo_perdedor:
            # Buscar o crear llave de tercer lugar
            llave_3er, _ = LlaveEliminatoria.objects.get_or_create(
                organizacion=organizacion,
                torneo=llave_db.torneo,
                fase='tercer_lugar',
                numero_llave=1,
                defaults={
                    'formato': 'partido_unico',
                    'estado': 'pendiente',
                    'orden_visual': 99
                }
            )
            # Semifinal 1 coloca local, Semifinal 2 coloca visitante
            if llave_db.numero_llave == 1:
                llave_3er.equipo_local = equipo_perdedor
            else:
                llave_3er.equipo_visitante = equipo_perdedor

            if llave_3er.equipo_local and llave_3er.equipo_visitante:
                llave_3er.estado = 'programada'
            llave_3er.save()

        # MANEJO DE FINAL Y CUADRO DE HONOR
        if llave_db.fase == 'final':
            campeon = equipo_ganador
            subcampeon = equipo_perdedor

            tercer_lugar = None
            cuarto_lugar = None

            if llave_db.torneo.disputar_tercer_lugar:
                llave_3er = LlaveEliminatoria.objects.filter(
                    torneo=llave_db.torneo,
                    fase='tercer_lugar',
                    estado='finalizada'
                ).first()
                if llave_3er and llave_3er.ganador:
                    tercer_lugar = llave_3er.ganador
                    if llave_3er.equipo_local and llave_3er.equipo_visitante:
                        cuarto_lugar = llave_3er.equipo_visitante if tercer_lugar == llave_3er.equipo_local else llave_3er.equipo_local

            ResultadoFinalTorneo.objects.update_or_create(
                torneo=llave_db.torneo,
                defaults={
                    'organizacion': organizacion,
                    'campeon': campeon,
                    'subcampeon': subcampeon,
                    'tercer_lugar': tercer_lugar,
                    'cuarto_lugar': cuarto_lugar,
                    'confirmado_por': usuario,
                    'observaciones': datos_confirmacion.get('observaciones', f"Campeón {campeon.nombre} confirmado tras la Final.")
                }
            )

        # Si se finaliza la llave de 3er lugar después de la final, actualizar el ResultadoFinalTorneo
        if llave_db.fase == 'tercer_lugar':
            res_final = ResultadoFinalTorneo.objects.filter(torneo=llave_db.torneo).first()
            if res_final:
                res_final.tercer_lugar = equipo_ganador
                if equipo_perdedor:
                    res_final.cuarto_lugar = equipo_perdedor
                res_final.save()

        # Auditoría
        detalles_audit = f"Confirmado ganador {equipo_ganador.nombre} en {llave_db.get_fase_display()} Llave #{llave_db.numero_llave} (Método: {llave_db.get_metodo_definicion_display()})."
        if es_decision_admin:
            detalles_audit += f" Motivo: {motivo_admin}"

        BitacoraTorneo.objects.create(
            organizacion=organizacion,
            torneo=llave_db.torneo,
            usuario=usuario,
            accion="Confirmación de Ganador de Llave",
            detalles=detalles_audit
        )

    return llave_db


def procesar_bye_llave(llave, usuario, organizacion):
    """
    Confirma el avance de un equipo con pase directo (BYE) sin requerir partidos.
    """
    ganador = llave.equipo_local or llave.equipo_visitante
    if not ganador:
        raise ValidationError("No existe un equipo asignado a esta llave BYE.")

    return confirmar_ganador_llave(
        llave=llave,
        ganador_id=ganador.id,
        usuario=usuario,
        organizacion=organizacion,
        datos_confirmacion={'motivo_admin': 'Pase Directo por BYE'}
    )


def reabrir_llave_eliminatoria(llave, usuario, organizacion, motivo):
    """
    Permite a un superadmin reabrir una llave eliminatoria finalizada si no han iniciado partidos posteriores.
    """
    if usuario.role != 'superadmin':
        raise ValidationError("Solo un superadmin puede reabrir una llave eliminatoria finalizada.")

    if not motivo or len(motivo.strip()) < 5:
        raise ValidationError("Debes especificar un motivo válido de al menos 5 caracteres para reabrir la llave.")

    with transaction.atomic():
        llave_db = LlaveEliminatoria.objects.select_for_update().get(id=llave.id, organizacion=organizacion)

        if llave_db.estado != 'finalizada':
            raise ValidationError("La llave especificada no está en estado finalizada.")

        # Verificar si existen partidos iniciados o finalizados en rondas posteriores que dependan de esta llave
        siguiente = llave_db.siguiente_llave
        if siguiente:
            partidos_posteriores = Partido.objects.filter(
                organizacion=organizacion,
                torneo=llave_db.torneo,
                equipo_local__in=[siguiente.equipo_local, siguiente.equipo_visitante],
                equipo_visitante__in=[siguiente.equipo_local, siguiente.equipo_visitante],
                estado__in=['en_curso', 'finalizado']
            ).count()

            if partidos_posteriores > 0:
                raise ValidationError("No se puede reabrir la llave porque ya existen partidos iniciados o finalizados en la ronda posterior.")

            # Retirar al ganador de la siguiente llave
            if llave_db.posicion_siguiente_llave == 'local':
                siguiente.equipo_local = None
            else:
                siguiente.equipo_visitante = None

            if siguiente.estado != 'finalizada':
                siguiente.estado = 'pendiente'
            siguiente.save()

            # Borrar partidos programados (no iniciados) de la siguiente llave
            Partido.objects.filter(
                organizacion=organizacion,
                torneo=llave_db.torneo,
                fase=siguiente.fase,
                equipo_local__in=[llave_db.ganador],
                estado='programado'
            ).delete()
            Partido.objects.filter(
                organizacion=organizacion,
                torneo=llave_db.torneo,
                fase=siguiente.fase,
                equipo_visitante__in=[llave_db.ganador],
                estado='programado'
            ).delete()

        # Si era la final o afectaba el cuadro de honor, borrar ResultadoFinalTorneo
        if llave_db.fase in ['final', 'tercer_lugar']:
            ResultadoFinalTorneo.objects.filter(torneo=llave_db.torneo).delete()

        # Resets de la llave actual
        ganador_previo = llave_db.ganador
        llave_db.ganador = None
        llave_db.estado = 'programada' if (llave_db.equipo_local and llave_db.equipo_visitante and not llave_db.es_bye) else 'pendiente'
        llave_db.definido_por_penales = False
        llave_db.penales_local = None
        llave_db.penales_visitante = None
        llave_db.hubo_prorroga = False
        llave_db.metodo_definicion = None
        llave_db.confirmado_por = None
        llave_db.fecha_confirmacion = None
        llave_db.save()

        BitacoraTorneo.objects.create(
            organizacion=organizacion,
            torneo=llave_db.torneo,
            usuario=usuario,
            accion="Reapertura de Llave Eliminatoria",
            detalles=f"Se reabrió {llave_db.get_fase_display()} Llave #{llave_db.numero_llave}. Ganador previo revocados ({ganador_previo.nombre if ganador_previo else 'N/A'}). Motivo: {motivo}"
        )

    return llave_db


def programar_partidos_ronda_siguiente(torneo, organizacion, usuario, datos):
    """
    Crea los partidos para las llaves preparadas de la ronda siguiente.
    """
    fase = datos.get('fase')
    fecha_inicial_str = datos.get('fecha_inicial', timezone.now().strftime('%Y-%m-%d'))
    hora_inicial_str = datos.get('hora_inicial', '14:00')
    intervalo_dias = int(datos.get('intervalo_dias', 7))
    estadio = datos.get('estadio', 'Estadio Principal')
    formato = datos.get('formato', 'partido_unico')

    try:
        dt_base = datetime.strptime(f"{fecha_inicial_str} {hora_inicial_str}", '%Y-%m-%d %H:%M')
        dt_base = timezone.make_aware(dt_base)
    except Exception:
        dt_base = timezone.now() + timedelta(days=1)

    llaves = list(
        LlaveEliminatoria.objects.filter(
            torneo=torneo,
            fase=fase,
            equipo_local__isnull=False,
            equipo_visitante__isnull=False,
            es_bye=False
        ).order_by('numero_llave')
    )

    if not llaves:
        raise ValidationError(f"No existen llaves preparadas con ambos equipos para la fase {fase}.")

    partidos_creados = []
    with transaction.atomic():
        for idx, ll in enumerate(llaves, start=1):
            # Eliminar partidos previa programación no iniciados de esta llave
            Partido.objects.filter(
                organizacion=organizacion,
                torneo=torneo,
                fase=fase,
                equipo_local=ll.equipo_local,
                equipo_visitante=ll.equipo_visitante,
                estado='programado'
            ).delete()

            dt_partido = dt_base + timedelta(hours=(idx - 1) * 2)

            p1 = Partido.objects.create(
                organizacion=organizacion,
                torneo=torneo,
                fase=fase,
                numero_vuelta=1,
                equipo_local=ll.equipo_local,
                equipo_visitante=ll.equipo_visitante,
                fecha_hora=dt_partido,
                estadio=estadio,
                jornada=idx,
                estado='programado',
                temporada=torneo.temporada
            )
            partidos_creados.append(p1)

            if formato == 'ida_vuelta':
                dt_vuelta = dt_base + timedelta(days=intervalo_dias, hours=(idx - 1) * 2)
                p2 = Partido.objects.create(
                    organizacion=organizacion,
                    torneo=torneo,
                    fase=fase,
                    numero_vuelta=2,
                    equipo_local=ll.equipo_visitante,
                    equipo_visitante=ll.equipo_local,
                    fecha_hora=dt_vuelta,
                    estadio=estadio,
                    jornada=idx,
                    estado='programado',
                    temporada=torneo.temporada
                )
                partidos_creados.append(p2)

            ll.formato = formato
            ll.estado = 'programada'
            ll.save()

        BitacoraTorneo.objects.create(
            organizacion=organizacion,
            torneo=torneo,
            usuario=usuario,
            accion="Programación de Ronda Eliminatoria",
            detalles=f"Programados {len(partidos_creados)} partidos para la fase {fase} (Formato: {formato})."
        )

    return partidos_creados


def generar_excel_cuadro_eliminatorio_completo(torneo, organizacion):
    """
    Genera un libro Excel (.xlsx) con la estructura completa del torneo, clasificatorios,
    todas las fases eliminatorias, tercer lugar, cuadro de honor (Campeón) y auditoría.
    """
    if not OPENPYXL_AVAILABLE:
        raise ImportError("openpyxl no está disponible.")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    gold_fill = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
    gold_font = Font(name="Calibri", size=11, bold=True, color="854D0E")
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    # 1. Hoja Clasificados
    ws_clas = wb.create_sheet(title="Clasificados")
    ws_clas.append([f"CLASIFICADOS DEFINITIVOS - {torneo.nombre.upper()}"])
    ws_clas.append(["BOMBO", "EQUIPO", "GRUPO", "POSICIÓN", "PTS", "DG", "GF", "FECHA CONFIRMACIÓN"])
    ws_clas['A1'].font = title_font

    for c in range(1, 9):
        ws_clas.cell(row=2, column=c).fill = header_fill
        ws_clas.cell(row=2, column=c).font = header_font

    clasificados = ClasificadoTorneo.objects.filter(torneo=torneo, activo=True).order_by('bombo', 'posicion_grupo')
    for item in clasificados:
        ws_clas.append([
            f"Bombo {item.bombo}",
            item.equipo.nombre,
            item.grupo.nombre,
            item.posicion_grupo,
            item.puntos,
            item.diferencia_goles,
            item.goles_favor,
            item.fecha_confirmacion.strftime('%Y-%m-%d %H:%M')
        ])
        for col_idx in range(1, 9):
            ws_clas.cell(row=ws_clas.max_row, column=col_idx).border = thin_border

    # 2. Hoja Cuadro Eliminatorio
    ws_cuadro = wb.create_sheet(title="Cuadro Eliminatorio")
    ws_cuadro.append([f"CUADRO ELIMINATORIO - {torneo.nombre.upper()}"])
    ws_cuadro.append(["FASE", "LLAVE #", "LOCAL", "VISITANTE", "GLOBAL", "PENALES", "GANADOR", "MÉTODO", "ESTADO"])
    ws_cuadro['A1'].font = title_font

    for c in range(1, 10):
        ws_cuadro.cell(row=2, column=c).fill = header_fill
        ws_cuadro.cell(row=2, column=c).font = header_font

    llaves = LlaveEliminatoria.objects.filter(torneo=torneo).order_by('fase', 'numero_llave')
    for ll in llaves:
        loc_name = ll.equipo_local.nombre if ll.equipo_local else ("BYE" if ll.es_bye else "Por Definir")
        vis_name = ll.equipo_visitante.nombre if ll.equipo_visitante else ("BYE" if ll.es_bye else "Por Definir")
        gan_name = ll.ganador.nombre if ll.ganador else ("BYE" if ll.es_bye else "-")
        global_str = f"{ll.marcador_global_local} - {ll.marcador_global_visitante}" if ll.estado in ['lista_para_confirmar', 'finalizada'] else "-"
        penales_str = f"{ll.penales_local} - {ll.penales_visitante}" if ll.definido_por_penales else "-"

        ws_cuadro.append([
            ll.get_fase_display(),
            ll.numero_llave,
            loc_name,
            vis_name,
            global_str,
            penales_str,
            gan_name,
            ll.get_metodo_definicion_display() if ll.metodo_definicion else "-",
            ll.get_estado_display()
        ])
        for col_idx in range(1, 10):
            ws_cuadro.cell(row=ws_cuadro.max_row, column=col_idx).border = thin_border

    # 3. Hoja Campeón y Cuadro de Honor
    ws_camp = wb.create_sheet(title="Cuadro de Honor")
    ws_camp.append([f"CUADRO DE HONOR Y CAMPEÓN - {torneo.nombre.upper()}"])
    ws_camp.append(["POSICIÓN", "EQUIPO", "CONFIRMADO POR", "FECHA"])
    ws_camp['A1'].font = title_font

    for c in range(1, 5):
        ws_camp.cell(row=2, column=c).fill = header_fill
        ws_camp.cell(row=2, column=c).font = header_font

    res_final = ResultadoFinalTorneo.objects.filter(torneo=torneo).first()
    if res_final:
        podio = [
            ("CAMPEÓN", res_final.campeon),
            ("SUBCAMPEÓN", res_final.subcampeon),
            ("TERCER LUGAR", res_final.tercer_lugar),
            ("CUARTO LUGAR", res_final.cuarto_lugar),
        ]
        for pos_label, eq in podio:
            if eq:
                ws_camp.append([
                    pos_label,
                    eq.nombre,
                    res_final.confirmado_por.get_full_name() if res_final.confirmado_por else "Sistema",
                    res_final.fecha_confirmacion.strftime('%Y-%m-%d %H:%M')
                ])
                curr_row = ws_camp.max_row
                for col_idx in range(1, 5):
                    ws_camp.cell(row=curr_row, column=col_idx).border = thin_border
                if pos_label == "CAMPEÓN":
                    ws_camp.cell(row=curr_row, column=2).fill = gold_fill
                    ws_camp.cell(row=curr_row, column=2).font = gold_font

    # 4. Hoja Auditoría
    ws_audit = wb.create_sheet(title="Auditoría")
    ws_audit.append([f"BITÁCORA DE AUDITORÍA - {torneo.nombre.upper()}"])
    ws_audit.append(["FECHA Y HORA", "USUARIO", "ACCIÓN", "DETALLES"])
    ws_audit['A1'].font = title_font

    for c in range(1, 5):
        ws_audit.cell(row=2, column=c).fill = header_fill
        ws_audit.cell(row=2, column=c).font = header_font

    logs = BitacoraTorneo.objects.filter(torneo=torneo).order_by('-fecha_hora')[:100]
    for lg in logs:
        ws_audit.append([
            lg.fecha_hora.strftime('%Y-%m-%d %H:%M'),
            lg.usuario.get_full_name() if lg.usuario else "Sistema",
            lg.accion,
            lg.detalles or ""
        ])
        for col_idx in range(1, 5):
            ws_audit.cell(row=ws_audit.max_row, column=col_idx).border = thin_border

    # Ajustar ancho de columnas
    for ws_curr in [ws_clas, ws_cuadro, ws_camp, ws_audit]:
        for col in ws_curr.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_curr.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
